"""R18 - PDF and Excel export for Admin, Faculty and Student, built from the same data and rules as the views.

Seed: F001 teaches 9 sessions (10 periods); each batch has 16 sessions (18 periods); 32 sessions in all. faculty1 = F001 (Dr. Rakesh Sharma),
student1/2 = batch A, student3/4 = batch B.
"""
import json
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import AcademicTerm, Role, Subject, TermType, Timetable, TimetableEntry, TimetableStatus
from app.modules.jiya import export_render
from app.tests.conftest import auth_header
from app.tests.jiya.test_generation import admin_id, disable_slots, entries, gen, term
from app.tests.jiya.test_validation import first_entry, refresh

URL = "/api/v1/timetables"
FAC, STU = "/api/v1/faculty/timetable/export", "/api/v1/student/timetable/export"
S = TimetableStatus
PDF, XLSX = "application/pdf", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def data(app, seeded):
    return None


def publish(client, admin, tid):
    for a in ("submit", "approve", "publish"):
        assert client.post(f"{URL}/{tid}/{a}", json={}, headers=admin).status_code == 200


@pytest.fixture()
def published(client, admin_headers):
    tid = gen()["id"]
    publish(client, admin_headers, tid)
    return tid


def as_(client, email):
    return auth_header(client, f"{email}@timetable.test")


def workbook(response):
    assert response.status_code == 200, response.get_data(as_text=True)[:300]
    return load_workbook(BytesIO(response.data))


def cell_texts(wb):
    return [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]


def class_rows(wb):
    ws = wb["Class list"]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


def check_pdf(response):
    assert response.status_code == 200 and response.mimetype == PDF
    data = response.data
    assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF")
    return data


class TestAdminExport:
    def test_excel_has_one_sheet_per_batch_and_a_class_list(self, client, admin_headers, data):
        tid = gen()["id"]
        r = client.get(f"{URL}/{tid}/export?format=xlsx", headers=admin_headers)
        assert r.mimetype == XLSX
        wb = workbook(r)
        assert wb.sheetnames == ["CSE Sem 3 - A", "CSE Sem 3 - B", "Class list"]
        assert len(class_rows(wb)) == 32                                  # every session of the term
        header = [c.value for c in wb["Class list"][1]]
        assert header == ["Section", "Day", "Start", "End", "Periods", "Duration", "Subject code", "Subject", "Teacher", "Room", "Type"]

    def test_excel_grid_has_days_times_a_break_and_a_merged_two_period_lab(self, client, admin_headers, data):
        tid = gen()["id"]
        ws = workbook(client.get(f"{URL}/{tid}/export?format=xlsx", headers=admin_headers))["CSE Sem 3 - A"]
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert {"Time", "Monday", "Friday", "Break", "09:00–10:00"} <= set(texts)
        assert any("Data Structures Lab" in t and "Lab · 2 periods · 2 h" in t for t in texts)
        vertical = [m for m in ws.merged_cells.ranges if m.min_col == m.max_col and m.max_row - m.min_row == 1]
        assert len(vertical) == 2                                        # batch A has two labs, each spans two period rows
        assert ws.freeze_panes is not None and ws.page_setup.orientation == "landscape"

    def test_pdf_is_valid_with_headings_and_the_draft_label(self, client, admin_headers, data):
        tid = gen()["id"]
        r = client.get(f"{URL}/{tid}/export?format=pdf", headers=admin_headers)
        pdf = check_pdf(r)
        assert pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") >= 2   # at least one page per batch
        for needle in (b"Timetable - 2026-27 Odd timetable v1", b"CSE Sem 3 - A", b"CSE Sem 3 - B", b"Data Structures Lab", b"NOT the published timetable", b"Class list"):
            assert needle in pdf
        assert "attachment" in r.headers["Content-Disposition"] and r.headers["Content-Disposition"].endswith('.pdf"')
        assert r.headers["Cache-Control"] == "no-store"

    def test_a_published_timetable_carries_no_draft_label(self, client, admin_headers, published):
        pdf = client.get(f"{URL}/{published}/export?format=pdf", headers=admin_headers).data
        assert b"NOT the published timetable" not in pdf and b"Status: Published" in pdf
        wb = workbook(client.get(f"{URL}/{published}/export?format=xlsx", headers=admin_headers))
        assert any("Status: Published" == t for t in cell_texts(wb)) and not any("NOT the published" in t for t in cell_texts(wb))

    @pytest.mark.parametrize("status", [S.DRAFT, S.VALIDATED, S.PENDING_APPROVAL, S.APPROVED, S.REJECTED, S.PUBLISHED, S.ARCHIVED])
    def test_every_status_can_be_exported_and_is_labelled(self, client, admin_headers, data, status):
        tid = gen()["id"]
        refresh(tid).status = status
        db.session.commit()
        for fmt in ("pdf", "xlsx"):
            r = client.get(f"{URL}/{tid}/export?format={fmt}", headers=admin_headers)
            assert r.status_code == 200
        pdf = client.get(f"{URL}/{tid}/export?format=pdf", headers=admin_headers).data
        assert (status.value.replace("_", " ").capitalize().encode() in pdf)

    def test_unscheduled_sessions_get_their_own_table(self, client, admin_headers, data):
        disable_slots(lambda s: s.day_of_week > 1)
        tid = gen()["id"]
        wb = workbook(client.get(f"{URL}/{tid}/export?format=xlsx", headers=admin_headers))
        assert "Unscheduled sessions" in wb.sheetnames
        rows = [[c.value for c in row] for row in wb["Unscheduled sessions"].iter_rows(min_row=2)]
        assert rows and all(len(r[3]) > 20 for r in rows)
        assert b"Unscheduled sessions" in client.get(f"{URL}/{tid}/export?format=pdf", headers=admin_headers).data

    def test_an_empty_timetable_still_gives_a_valid_file_with_a_note(self, client, admin_headers, data):
        t = Timetable(term_id=term().id, name="Empty one", version=1, generated_by=admin_id())
        db.session.add(t)
        db.session.commit()
        assert b"No classes are scheduled in this timetable." in check_pdf(client.get(f"{URL}/{t.id}/export?format=pdf", headers=admin_headers))
        wb = workbook(client.get(f"{URL}/{t.id}/export?format=xlsx", headers=admin_headers))
        assert any("No classes are scheduled" in x for x in cell_texts(wb)) and class_rows(wb) == []

    def test_text_that_looks_like_a_formula_stays_text(self, client, admin_headers, data):
        tid = gen()["id"]
        subject = first_entry(tid, "CS301", "A").assignment.subject
        subject.name = '=HYPERLINK("http://evil.example","x")'
        db.session.commit()
        wb = workbook(client.get(f"{URL}/{tid}/export?format=xlsx", headers=admin_headers))
        hits = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
        assert hits and all(c.data_type == "s" for c in hits)
        check_pdf(client.get(f"{URL}/{tid}/export?format=pdf", headers=admin_headers))

    def test_markup_characters_do_not_break_the_pdf(self, client, admin_headers, data):
        tid = gen()["id"]
        subject = first_entry(tid, "CS301", "A").assignment.subject
        subject.name = "Maths <b>& Logic</i> > 2"
        db.session.commit()
        assert b"Maths" in check_pdf(client.get(f"{URL}/{tid}/export?format=pdf", headers=admin_headers))

    def test_invalid_requests(self, client, admin_headers, data):
        tid = gen()["id"]
        for q in ("", "?format=", "?format=docx", "?format=PDFX", "?format=pdf,xlsx"):
            r = client.get(f"{URL}/{tid}/export{q}", headers=admin_headers)
            assert r.status_code == 400 and r.get_json()["error"]["details"]["format"], q
        for q in ("?format=pdf&batch_id=2", "?format=pdf&faculty_id=1", "?format=pdf&term_id=1"):
            assert client.get(f"{URL}/{tid}/export{q}", headers=admin_headers).status_code == 400
        assert client.get(f"{URL}/9999/export?format=pdf", headers=admin_headers).status_code == 404

    def test_only_get_exists(self, client, admin_headers, data):
        tid = gen()["id"]
        assert client.post(f"{URL}/{tid}/export?format=pdf", json={}, headers=admin_headers).status_code == 405

    def test_exporting_changes_nothing(self, client, admin_headers, data):
        tid = gen()["id"]
        before = (refresh(tid).status, refresh(tid).validation_passed, sorted(e.id for e in entries(tid)))
        client.get(f"{URL}/{tid}/export?format=pdf", headers=admin_headers)
        client.get(f"{URL}/{tid}/export?format=xlsx", headers=admin_headers)
        assert (refresh(tid).status, refresh(tid).validation_passed, sorted(e.id for e in entries(tid))) == before

    def test_faculty_and_students_cannot_use_the_admin_export(self, client, published, faculty_headers, student_headers):
        for h in (faculty_headers, student_headers):
            assert client.get(f"{URL}/{published}/export?format=pdf", headers=h).status_code == 403
        assert client.get(f"{URL}/{published}/export?format=pdf").status_code == 401


class TestFacultyExport:
    def test_own_classes_only_in_excel(self, client, published):
        wb = workbook(client.get(f"{FAC}?format=xlsx", headers=as_(client, "faculty1")))
        assert wb.sheetnames == ["Weekly timetable", "Class list"]
        rows = class_rows(wb)
        assert len(rows) == 9 and sum(r[3] for r in rows) == 10 and {r[5] for r in rows} == {"CS301", "CS351"}
        assert {r[7] for r in rows} == {"CSE Sem 3 - A", "CSE Sem 3 - B"}      # the batch column for a teacher
        texts = " ".join(cell_texts(wb))
        for other in ("Meera Iyer", "Sanjay Gupta", "Kavita Rao", "CS302", "CS303", "MA301"):
            assert other not in texts
        assert "Dr. Rakesh Sharma (F001)" in texts and "Version 1" in texts

    def test_pdf_is_valid_and_private(self, client, published):
        r = client.get(f"{FAC}?format=pdf", headers=as_(client, "faculty1"))
        pdf = check_pdf(r)
        assert b"My teaching timetable" in pdf and b"Dr. Rakesh Sharma" in pdf and b"Data Structures Lab" in pdf
        for leak in (b"Meera Iyer", b"CS303", b"@timetable", b"password", b"validation"):
            assert leak not in pdf
        assert "F001" in r.headers["Content-Disposition"]

    def test_another_teacher_gets_a_different_file(self, client, published):
        f5 = workbook(client.get(f"{FAC}?format=xlsx", headers=as_(client, "faculty5")))
        assert [r[5] for r in class_rows(f5)] == ["CS303"] * 3

    def test_a_teacher_without_classes_gets_a_valid_file_with_a_note(self, client, published):
        from app.models import Faculty, User
        from app.tests.conftest import login, make_user
        make_user("loner@timetable.test", Role.FACULTY, "Password123")
        user = db.session.scalar(db.select(User).filter_by(email="loner@timetable.test"))
        db.session.add(Faculty(user_id=user.id, employee_code="F900", department_id=db.session.scalar(db.select(Faculty.department_id)), max_weekly_hours=10))
        db.session.commit()
        h = {"Authorization": f"Bearer {login(client, 'loner@timetable.test', 'Password123').get_json()['data']['token']}"}
        assert b"You have no classes in this timetable." in check_pdf(client.get(f"{FAC}?format=pdf", headers=h))
        wb = workbook(client.get(f"{FAC}?format=xlsx", headers=h))
        assert any("no classes" in t for t in cell_texts(wb)) and class_rows(wb) == []

    def test_requests_for_someone_elses_data_are_refused(self, client, published):
        h = as_(client, "faculty1")
        for q in ("faculty_id=5", "employee_code=F005", "batch_id=1", "user_id=3", "timetable_id=1", "student_id=1"):
            r = client.get(f"{FAC}?format=pdf&{q}", headers=h)
            assert r.status_code == 400 and r.mimetype == "application/json", q

    def test_roles_are_separated(self, client, published, admin_headers):
        assert client.get(f"{FAC}?format=pdf", headers=as_(client, "student1")).status_code == 403
        assert client.get(f"{FAC}?format=pdf", headers=admin_headers).status_code == 403
        assert client.get(f"{FAC}?format=pdf").status_code == 401

    def test_a_faculty_account_without_a_record(self, client, published):
        from app.tests.conftest import login, make_user
        make_user("norec@timetable.test", Role.FACULTY, "Password123")
        h = {"Authorization": f"Bearer {login(client, 'norec@timetable.test', 'Password123').get_json()['data']['token']}"}
        r = client.get(f"{FAC}?format=pdf", headers=h)
        assert r.status_code == 409 and r.get_json()["error"]["code"] == "no_profile"


class TestStudentExport:
    def test_own_batch_only_in_excel(self, client, published):
        a = workbook(client.get(f"{STU}?format=xlsx", headers=as_(client, "student1")))
        b = workbook(client.get(f"{STU}?format=xlsx", headers=as_(client, "student3")))
        assert a.sheetnames == ["CSE Sem 3 - A", "Class list"] and b.sheetnames == ["CSE Sem 3 - B", "Class list"]
        assert len(class_rows(a)) == 16 == len(class_rows(b)) and sum(r[3] for r in class_rows(a)) == 18
        assert "CSE Sem 3 - B" not in " ".join(cell_texts(a)) and "CSE Sem 3 - A" not in " ".join(cell_texts(b))
        assert [c.value for c in a["Class list"][1]][7] == "Teacher"
        assert {r[7] for r in class_rows(a)} != {r[7] for r in class_rows(b)}

    def test_pdf_is_valid_and_has_only_my_batch(self, client, published):
        r = client.get(f"{STU}?format=pdf", headers=as_(client, "student1"))
        pdf = check_pdf(r)
        assert b"Class timetable" in pdf and b"CSE Sem 3 - A" in pdf and b"Aarav Singh" in pdf and b"CSE23001" not in pdf
        for leak in (b"CSE Sem 3 - B", b"@timetable", b"password"):
            assert leak not in pdf
        assert "CSE-Sem-3-A" in r.headers["Content-Disposition"]

    def test_requests_for_another_batch_are_refused(self, client, published):
        h = as_(client, "student1")
        for q in ("batch_id=2", "batch=B", "student_id=3", "user_id=4", "timetable_id=1", "faculty_id=2"):
            assert client.get(f"{STU}?format=xlsx&{q}", headers=h).status_code == 400, q

    def test_roles_are_separated(self, client, published, admin_headers):
        assert client.get(f"{STU}?format=pdf", headers=as_(client, "faculty1")).status_code == 403
        assert client.get(f"{STU}?format=pdf", headers=admin_headers).status_code == 403
        assert client.get(f"{STU}?format=pdf").status_code == 401

    def test_only_get_exists(self, client, published):
        assert client.post(f"{STU}?format=pdf", json={}, headers=as_(client, "student1")).status_code == 405
        assert client.delete(f"{STU}?format=pdf", headers=as_(client, "student1")).status_code == 405


@pytest.mark.parametrize("path,who", [(FAC, "faculty1"), (STU, "student1")])
class TestNothingUnpublishedIsExported:
    @pytest.mark.parametrize("status", [S.DRAFT, S.VALIDATED, S.PENDING_APPROVAL, S.APPROVED, S.REJECTED, S.ARCHIVED])
    def test_no_file_for_a_non_published_status(self, client, data, path, who, status):
        tid = gen()["id"]
        refresh(tid).status = status
        db.session.commit()
        for fmt in ("pdf", "xlsx"):
            r = client.get(f"{path}?format={fmt}", headers=as_(client, who))
            assert r.status_code == 404 and r.get_json()["error"]["code"] == "no_published_timetable"
            assert "2026-27 Odd" in r.get_json()["error"]["message"]

    def test_no_timetable_at_all(self, client, data, path, who):
        r = client.get(f"{path}?format=pdf", headers=as_(client, who))
        assert r.status_code == 404 and r.mimetype == "application/json"

    def test_edits_to_a_draft_never_reach_the_export(self, client, admin_headers, published, path, who):
        before = client.get(f"{path}?format=xlsx", headers=as_(client, who)).data
        draft = gen()["id"]
        e = first_entry(draft, "CS301", "A")
        opts = client.get(f"{URL}/{draft}/entries/{e.id}/edit-options", headers=admin_headers).get_json()["data"]["options"]
        t = next(o for o in opts if not o["is_current"])
        client.patch(f"{URL}/{draft}/entries/{e.id}", json={"time_slot_id": t["time_slot_id"], "room_id": t["rooms"][0]["id"]}, headers=admin_headers)
        after = client.get(f"{path}?format=xlsx", headers=as_(client, who)).data
        assert cell_texts(load_workbook(BytesIO(before))) == cell_texts(load_workbook(BytesIO(after)))

    def test_the_superseded_version_is_not_exported(self, client, admin_headers, published, path, who):
        v2 = gen()["id"]
        publish(client, admin_headers, v2)
        r = client.get(f"{path}?format=pdf", headers=as_(client, who))
        assert b"Version 2" in r.data and b"Version 1" not in r.data and "v2" in r.headers["Content-Disposition"]
        assert any("Version 2" in t for t in cell_texts(workbook(client.get(f"{path}?format=xlsx", headers=as_(client, who)))))

    @pytest.mark.parametrize("q", ["", "&format=", "?x"])
    def test_a_missing_or_bad_format_is_a_clear_400(self, client, published, path, who, q):
        r = client.get(f"{path}{q if q.startswith('?') else '?' + q.lstrip('&') if q else ''}", headers=as_(client, who))
        assert r.status_code == 400 and r.get_json()["success"] is False

    def test_unsupported_formats(self, client, published, path, who):
        for fmt in ("csv", "docx", "html", "pdf%00"):
            assert client.get(f"{path}?format={fmt}", headers=as_(client, who)).status_code == 400


class TestTerms:
    def _other_term_with_a_class(self, published_tid):
        t = AcademicTerm(name="Spring 2027", academic_year="2026-27", term_type=TermType.EVEN, start_date=date(2027, 1, 1), end_date=date(2027, 5, 1), is_active=False)
        db.session.add(t)
        db.session.commit()
        tt = Timetable(term_id=t.id, name="Spring v1", version=1, status=S.PUBLISHED, generated_by=admin_id())
        db.session.add(tt)
        db.session.commit()
        src = first_entry(published_tid, "CS301", "A")
        db.session.add(TimetableEntry(timetable_id=tt.id, assignment_id=src.assignment_id, faculty_id=src.faculty_id, batch_id=src.batch_id, room_id=src.room_id, time_slot_id=src.time_slot_id))
        db.session.commit()
        return t

    def test_the_chosen_term_decides_the_file(self, client, published):
        other = self._other_term_with_a_class(published)
        for path, who, active_rows in ((STU, "student1", 16), (FAC, "faculty1", 9)):
            h = as_(client, who)
            assert len(class_rows(workbook(client.get(f"{path}?format=xlsx", headers=h)))) == active_rows
            chosen = workbook(client.get(f"{path}?format=xlsx&term_id={other.id}", headers=h))
            assert len(class_rows(chosen)) == 1 and any("Spring 2027" in t for t in cell_texts(chosen))
            assert "Spring-2027" in client.get(f"{path}?format=pdf&term_id={other.id}", headers=h).headers["Content-Disposition"]

    def test_a_term_without_a_published_timetable(self, client, published):
        t = AcademicTerm(name="Winter 2027", academic_year="2026-27", term_type=TermType.EVEN, start_date=date(2027, 1, 1), end_date=date(2027, 5, 1), is_active=False)
        db.session.add(t)
        db.session.commit()
        r = client.get(f"{STU}?format=pdf&term_id={t.id}", headers=as_(client, "student1"))
        assert r.status_code == 404 and r.get_json()["error"]["code"] == "no_published_timetable" and "Winter 2027" in r.get_json()["error"]["message"]

    def test_an_unknown_term(self, client, published):
        r = client.get(f"{FAC}?format=pdf&term_id=9999", headers=as_(client, "faculty1"))
        assert r.status_code == 404 and r.get_json()["error"]["code"] == "term_not_found"

    @pytest.mark.parametrize("raw", ["x", "0", "-4"])
    def test_a_bad_term_value(self, client, published, raw):
        assert client.get(f"{STU}?format=pdf&term_id={raw}", headers=as_(client, "student1")).status_code == 400

    def test_the_active_term_without_a_published_timetable_does_not_fall_back_to_an_old_one(self, client, data):
        other = self._other_term_with_a_class(gen()["id"])
        r = client.get(f"{STU}?format=pdf", headers=as_(client, "student1"))
        assert r.status_code == 404 and r.get_json()["error"]["code"] == "no_published_timetable"
        assert client.get(f"{STU}?format=pdf&term_id={other.id}", headers=as_(client, "student1")).status_code == 200


class TestRenderers:
    def _doc(self, sessions=None, **kw):
        slot = lambda p, s, e, t="class": {"day_of_week": 0, "day_name": "Monday", "period_no": p, "start_time": s, "end_time": e, "slot_type": t}
        sessions = sessions if sessions is not None else [{"day_of_week": 0, "day_name": "Monday", "period_no": 1, "periods": 1, "start_time": "09:00", "end_time": "10:00",
                                                           "subject": {"code": "X1", "name": "Name", "class_type": "theory"}, "room": {"code": "R1", "name": "Room"}, "batch": {"name": "B"}}]
        return export_render.Document("T", ["Line"], [slot(1, "09:00", "10:00"), slot(2, "10:00", "11:00"), slot(3, "11:00", "12:00", "break")], [export_render.Section("S", sessions, kw.get("note"))])

    def test_duration_text(self):
        d = lambda start, end, n: export_render.duration({"start_time": start, "end_time": end, "periods": n})
        assert d("09:00", "10:00", 1) == "1 period · 1 h" and d("09:00", "11:00", 2) == "2 periods · 2 h" and d("09:00", "09:45", 1) == "1 period · 45 min" and d("09:00", "10:30", 1) == "1 period · 1 h 30 min"

    def test_sheet_names_are_valid_and_unique(self):
        used = set()
        names = [export_render.sheet_name(n, used) for n in ("A/B", "A/B", "x" * 50, "[bad]:name?", "")]
        assert len(set(n.lower() for n in names)) == 5 and all(len(n) <= 31 and not set(n) & set('[]:*?/\\') for n in names)

    def test_both_renderers_handle_no_sessions_and_no_slots(self):
        for doc in (self._doc(sessions=[]), export_render.Document("T", [], [], [export_render.Section("S", [])])):
            assert export_render.render_pdf(doc).startswith(b"%PDF") and load_workbook(BytesIO(export_render.render_xlsx(doc)))

    def test_a_non_latin_name_does_not_crash(self):
        s = [{"day_of_week": 0, "day_name": "Monday", "period_no": 1, "periods": 1, "start_time": "09:00", "end_time": "10:00",
              "subject": {"code": "H1", "name": "हिन्दी साहित्य", "class_type": "theory"}, "room": {"code": "R1", "name": "R"}, "batch": {"name": "B"}}]
        doc = self._doc(sessions=s)
        assert export_render.render_pdf(doc).startswith(b"%PDF")
        assert "हिन्दी साहित्य" in " ".join(cell_texts(load_workbook(BytesIO(export_render.render_xlsx(doc)))))
