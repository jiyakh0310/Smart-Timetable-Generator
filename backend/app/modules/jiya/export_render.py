"""Jiya - R18 renderers: turn one neutral `Document` into a PDF (reportlab) or an Excel workbook (openpyxl).

Nothing here knows about users, permissions or the database: it only lays out data that `export_service` already filtered.
A Document is a title, a few header lines, and sections; each section is a weekly grid (days x periods, labs spanning their
periods) followed by a class list with day, time, subject, teacher or batch, room and duration.
"""
from dataclasses import dataclass, field
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INDIGO, SOFT, LAB_TINT, BREAK_TINT, GRID = "4F46E5", "EEF2FF", "E0F2FE", "F1F5F9", "CBD5E1"


@dataclass
class Section:
    heading: str
    sessions: list[dict]
    # "student" sessions carry the teacher, "faculty" sessions carry the batch; the class list adapts to which one is present
    note: str | None = None


@dataclass
class Document:
    title: str
    lines: list[str]
    slots: list[dict]
    sections: list[Section]
    footer: str = ""
    extra_tables: list[tuple[str, list[str], list[list[str]]]] = field(default_factory=list)  # (heading, header row, rows) after the sections


# ---------- shared helpers ----------
def duration(session: dict) -> str:
    sh, sm = map(int, session["start_time"].split(":"))
    eh, em = map(int, session["end_time"].split(":"))
    minutes = (eh * 60 + em) - (sh * 60 + sm)
    h, m = divmod(minutes, 60)
    text = f"{h} h" if not m else (f"{h} h {m} min" if h else f"{m} min")
    return f"{session['periods']} period{'s' if session['periods'] != 1 else ''} · {text}"


def who(session: dict) -> str:
    return session["faculty"]["name"] if "faculty" in session else session["batch"]["name"] if "batch" in session else ""


def who_label(sections: list[Section]) -> str:
    first = next((s for sec in sections for s in sec.sessions), None)
    return "Batch" if first is not None and "batch" in first else "Teacher"


def week_structure(slots: list[dict]):
    days = sorted({s["day_of_week"]: s["day_name"] for s in slots}.items())
    periods = {}
    for s in sorted(slots, key=lambda s: (s["period_no"], s["day_of_week"])):
        periods.setdefault(s["period_no"], {"no": s["period_no"], "start": s["start_time"], "end": s["end_time"], "kind": s["slot_type"]})
    return days, [periods[k] for k in sorted(periods)]


def grid_cells(sessions: list[dict], day_keys: set, period_keys: set):
    """{(day, first_period): session}; sessions on a period the week no longer has are left for the class list only."""
    return {(s["day_of_week"], s["period_no"]): s for s in sessions if s["day_of_week"] in day_keys and s["period_no"] in period_keys}


def cell_lines(session: dict) -> list[str]:
    return [f"{session['subject']['code']} · {session['subject']['name']}", f"Room {session['room']['code']}", who(session)] + (["Lab · " + duration(session)] if session["subject"]["class_type"] == "lab" else [])


def sheet_name(text: str, used: set) -> str:
    clean = "".join("-" if ch in '[]:*?/\\' else ch for ch in text).strip() or "Sheet"
    base, n = clean[:31], 1
    while base.lower() in used:
        n += 1
        base = f"{clean[:28]}-{n}"
    used.add(base.lower())
    return base


# ---------- Excel ----------
def _put(cell, value, **style):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # always text: a name that starts with "=" must never become a formula
    for k, v in style.items():
        setattr(cell, k, v)


def render_xlsx(doc: Document) -> bytes:
    wb = Workbook()
    wb.properties.title = doc.title
    wb.properties.creator = "Smart Timetable System"
    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill, soft, lab_fill, break_fill = (PatternFill("solid", fgColor=c) for c in (INDIGO, SOFT, LAB_TINT, BREAK_TINT))
    wrap = Alignment(wrap_text=True, vertical="top")
    days, periods = week_structure(doc.slots)
    used: set = set()
    first = True
    for section in doc.sections:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet_name(section.heading, used)
        width = max(len(days) + 1, 4)
        _put(ws.cell(1, 1), doc.title, font=Font(size=15, bold=True, color=INDIGO))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        row = 2
        for line in [section.heading] + doc.lines:
            _put(ws.cell(row, 1), line, font=Font(bold=(line == section.heading), size=11 if line == section.heading else 10))
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            row += 1
        row += 1
        if section.note:
            _put(ws.cell(row, 1), section.note, font=Font(italic=True), alignment=wrap)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            ws.row_dimensions[row].height = 30
            row += 2
        if days and periods:
            header_row = row
            for c, text in enumerate(["Time"] + [d[1] for d in days], start=1):
                _put(ws.cell(row, c), text, font=Font(bold=True, color="FFFFFF"), fill=head_fill, alignment=Alignment(horizontal="center", vertical="center"), border=border)
            row += 1
            cells = grid_cells(section.sessions, {d[0] for d in days}, {p["no"] for p in periods})
            last_period_row = header_row + len(periods)
            for p in periods:
                _put(ws.cell(row, 1), f"{p['start']}–{p['end']}", font=Font(bold=True), fill=soft, alignment=Alignment(vertical="top"), border=border)
                if p["kind"] == "break":
                    for c in range(2, len(days) + 2):
                        _put(ws.cell(row, c), "Break" if c == 2 else None, fill=break_fill, alignment=Alignment(horizontal="center", vertical="center"), border=border)
                    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(days) + 1)
                    ws.row_dimensions[row].height = 18
                else:
                    ws.row_dimensions[row].height = 62
                    for c, (day_no, _) in enumerate(days, start=2):
                        cell = ws.cell(row, c)
                        cell.border = border
                        s = cells.get((day_no, p["no"]))
                        if s:
                            _put(cell, "\n".join(cell_lines(s)), alignment=wrap, fill=lab_fill if s["subject"]["class_type"] == "lab" else PatternFill())
                            if s["periods"] > 1:
                                last = min(row + s["periods"] - 1, last_period_row)
                                ws.merge_cells(start_row=row, start_column=c, end_row=last, end_column=c)
                                for rr in range(row, last + 1):
                                    ws.cell(rr, c).border = border
                row += 1
            ws.freeze_panes = ws.cell(header_row + 1, 2)
            ws.column_dimensions["A"].width = 14
            for c in range(2, len(days) + 2):
                ws.column_dimensions[get_column_letter(c)].width = 26
        elif not section.sessions:
            _put(ws.cell(row, 1), "No classes are scheduled.", font=Font(italic=True))
        _page_setup(ws, landscape_=True)

    # one flat, filter-friendly list of every class
    lw = wb.create_sheet(sheet_name("Class list", used))
    label = who_label(doc.sections)
    multi = len(doc.sections) > 1
    headers = (["Section"] if multi else []) + ["Day", "Start", "End", "Periods", "Duration", "Subject code", "Subject", label, "Room", "Type"]
    for c, text in enumerate(headers, start=1):
        _put(lw.cell(1, c), text, font=Font(bold=True, color="FFFFFF"), fill=head_fill, alignment=Alignment(horizontal="center"), border=border)
    r = 2
    for section in doc.sections:
        for s in section.sessions:
            values = ([section.heading] if multi else []) + [s["day_name"], s["start_time"], s["end_time"], s["periods"], duration(s).split(" · ")[1], s["subject"]["code"], s["subject"]["name"], who(s), s["room"]["code"],
                                                              "Lab" if s["subject"]["class_type"] == "lab" else "Theory"]
            for c, v in enumerate(values, start=1):
                _put(lw.cell(r, c), v, alignment=Alignment(vertical="top"), border=border)
            r += 1
    for i, w in enumerate(([22] if multi else []) + [12, 8, 8, 9, 12, 13, 32, 26, 10, 9], start=1):
        lw.column_dimensions[get_column_letter(i)].width = w
    lw.freeze_panes = "A2"
    if r > 2:
        lw.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r - 1}"
    _page_setup(lw, landscape_=True)

    for heading, header_row_, rows in doc.extra_tables:
        ew = wb.create_sheet(sheet_name(heading, used))
        for c, text in enumerate(header_row_, start=1):
            _put(ew.cell(1, c), text, font=Font(bold=True, color="FFFFFF"), fill=head_fill, border=border)
            ew.column_dimensions[get_column_letter(c)].width = 28 if c == len(header_row_) else 18
        for i, values in enumerate(rows, start=2):
            for c, v in enumerate(values, start=1):
                _put(ew.cell(i, c), v, alignment=wrap, border=border)
        ew.freeze_panes = "A2"
        _page_setup(ew, landscape_=True)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _page_setup(ws, landscape_=True):
    ws.page_setup.orientation = "landscape" if landscape_ else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Page &P of &N"


# ---------- PDF ----------
def _para(text, style):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)


def render_pdf(doc: Document) -> bytes:
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#" + INDIGO), alignment=0, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=6, spaceAfter=4)
    meta = ParagraphStyle("m", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#475569"), leading=12)
    cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=7.5, leading=9)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
    head = ParagraphStyle("hd", parent=cell, fontName="Helvetica-Bold", textColor=colors.white, alignment=1)
    note = ParagraphStyle("n", parent=styles["Italic"], fontSize=9)

    buffer = BytesIO()
    page = landscape(A4)
    pdf = SimpleDocTemplate(buffer, pagesize=page, leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.4 * cm, bottomMargin=1.5 * cm, title=doc.title, author="Smart Timetable System", pageCompression=0)
    usable = page[0] - 2.4 * cm
    days, periods = week_structure(doc.slots)
    label = who_label(doc.sections)

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(1.2 * cm, 0.8 * cm, (doc.footer or doc.title)[:120])
        canvas.drawRightString(page[0] - 1.2 * cm, 0.8 * cm, f"Page {d.page}")
        canvas.restoreState()

    story = []
    for index, section in enumerate(doc.sections):
        if index:
            story.append(PageBreak())
        story += [Paragraph(escape(doc.title), title), Paragraph(f"<b>{escape(section.heading)}</b>", h2)]
        story += [Paragraph(escape(line), meta) for line in doc.lines]
        story.append(Spacer(1, 6))
        if section.note:
            story += [Paragraph(escape(section.note), note), Spacer(1, 6)]
        if days and periods:
            cells = grid_cells(section.sessions, {d[0] for d in days}, {p["no"] for p in periods})
            data = [[_para("Time", head)] + [_para(d[1], head) for d in days]]
            style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + INDIGO)), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#" + GRID)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#" + SOFT)), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
            for r, p in enumerate(periods, start=1):
                row = [_para(f"{p['start']}–{p['end']}", cellb)]
                if p["kind"] == "break":
                    row += [_para("Break", ParagraphStyle("b", parent=cell, alignment=1))] + [""] * (len(days) - 1)
                    style += [("SPAN", (1, r), (-1, r)), ("BACKGROUND", (1, r), (-1, r), colors.HexColor("#" + BREAK_TINT))]
                else:
                    for c, (day_no, _) in enumerate(days, start=1):
                        s = cells.get((day_no, p["no"]))
                        if s:
                            lines = cell_lines(s)
                            row.append(Paragraph(f"<b>{escape(lines[0])}</b><br/>" + "<br/>".join(escape(x) for x in lines[1:]), cell))
                            if s["periods"] > 1 and r + s["periods"] - 1 <= len(periods):
                                style += [("SPAN", (c, r), (c, r + s["periods"] - 1))]
                            if s["subject"]["class_type"] == "lab":
                                style += [("BACKGROUND", (c, r), (c, min(r + s["periods"] - 1, len(periods))), colors.HexColor("#" + LAB_TINT))]
                        else:
                            row.append("")
                data.append(row)
            col = [2.3 * cm] + [(usable - 2.3 * cm) / len(days)] * len(days)
            story.append(Table(data, colWidths=col, style=TableStyle(style), repeatRows=1))
        if section.sessions:
            story += [Spacer(1, 10), Paragraph("Class list", h2)]
            rows = [[_para(h, head) for h in ("Day", "Time", "Subject", label, "Room", "Duration")]]
            for s in section.sessions:
                rows.append([_para(s["day_name"], cell), _para(f"{s['start_time']}–{s['end_time']}", cell), _para(f"{s['subject']['code']} · {s['subject']['name']}", cell), _para(who(s), cell), _para(s["room"]["code"], cell),
                             _para(("Lab · " if s["subject"]["class_type"] == "lab" else "") + duration(s), cell)])
            widths = [2.4 * cm, 2.6 * cm, usable - 2.4 * cm - 2.6 * cm - 5.2 * cm - 2.2 * cm - 4.4 * cm, 5.2 * cm, 2.2 * cm, 4.4 * cm]
            story.append(Table(rows, colWidths=widths, repeatRows=1, style=TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + INDIGO)), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#" + GRID)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])])))
        elif not section.note:
            story.append(Paragraph("No classes are scheduled.", note))
    for heading, header_row, rows in doc.extra_tables:
        story += [PageBreak(), Paragraph(escape(doc.title), title), Paragraph(f"<b>{escape(heading)}</b>", h2)]
        data = [[_para(h, head) for h in header_row]] + [[_para(v, cell) for v in r] for r in rows]
        story.append(Table(data, repeatRows=1, colWidths=[usable / len(header_row)] * len(header_row), style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + INDIGO)), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#" + GRID)), ("VALIGN", (0, 0), (-1, -1), "TOP")])))
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
